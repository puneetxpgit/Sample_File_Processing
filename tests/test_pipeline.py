import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from processor.classify_email import classify_email
from processor.export_kirke import assign_mop_names
from processor.pipeline import process_package
from processor.read_inputs import read_mop, read_spreadsheet
from processor.settings import load_settings
from processor.sow_catalog import SowCatalog
from processor.sow_model import SowClassifier

ROOT = Path(__file__).resolve().parents[1]
TOWER = """Please process.

-----Original Message-----
From: Jane GC <jane.gc@vendor.com>
Sent: Tuesday, October 20, 2026 8:14 AM
To: Tower MOP <tower-mop@example.com>
Subject: GILLETT

Schedule attached.
"""
GROUND = TOWER.replace("tower-mop@example.com", "ground-mop@example.com").replace("Tower MOP", "Ground MOP")


class ClassifyEmailTests(unittest.TestCase):
    def test_tower_ground_and_review(self):
        tower = classify_email({"id": "1", "body": TOWER}, ["tower-mop@example.com"], ["ground-mop@example.com"])
        self.assertEqual(tower["category"], "Tower")
        self.assertEqual(tower["status"], "classified")
        self.assertEqual(tower["originalFrom"], "jane.gc@vendor.com")
        self.assertIn("October 20, 2026", tower["originalDate"])
        self.assertEqual(tower["originalTo"], "tower-mop@example.com")

        ground = classify_email({"id": "2", "body": GROUND}, ["tower-mop@example.com"], ["ground-mop@example.com"])
        self.assertEqual(ground["category"], "Ground")

        both = TOWER.replace(
            "To: Tower MOP <tower-mop@example.com>",
            "To: Tower MOP <tower-mop@example.com>; Ground MOP <ground-mop@example.com>",
        )
        mixed = classify_email({"id": "3", "body": both}, ["tower-mop@example.com"], ["ground-mop@example.com"])
        self.assertEqual(mixed["category"], "Review")
        self.assertIn("both", mixed["reason"])

        missing = classify_email({"id": "4", "body": "No headers here, just a note."}, ["tower-mop@example.com"], ["ground-mop@example.com"])
        self.assertEqual(missing["category"], "Review")
        self.assertEqual(missing["status"], "review_required")

        empty = classify_email({"id": "5", "body": "   "}, ["tower-mop@example.com"], ["ground-mop@example.com"])
        self.assertEqual(empty["reason"], "Email body is missing or empty")

        unknown = classify_email({"id": "6", "body": TOWER}, ["other@example.com"], ["ground-mop@example.com"])
        self.assertEqual(unknown["category"], "Review")
        self.assertIn("Unrecognized", unknown["reason"])


class CatalogTests(unittest.TestCase):
    def test_catalog_categories_known_sows(self):
        settings = load_settings(ROOT)
        catalog = SowCatalog.load(settings.path(settings.catalog_path))
        tower_types = {example.request_type for example in catalog.examples if example.category == "Tower"}
        ground_types = {example.request_type for example in catalog.examples if example.category == "Ground"}
        self.assertIn("SNAP", tower_types)
        self.assertIn("Cdu to uADPF", tower_types)
        self.assertIn("Ground Prep", ground_types)
        self.assertNotIn("N/A", ground_types)

        pdu = catalog.match("Ground", "PDU upgrade", settings.request_types["Ground"])
        self.assertEqual(pdu["requestType"], "PDU")
        self.assertGreaterEqual(pdu["confidence"], 90)

        cdu = catalog.match("Tower", "CDU to uADPF Conversion", settings.request_types["Tower"])
        self.assertEqual(cdu["requestType"], "Cdu to uADPF")

        jerry = read_mop(ROOT / "JERRY_BOYLES_RD_MOD_PROJECT_MOP_2026.docx")
        ground_cdu = catalog.match("Ground", jerry.scope_text, settings.request_types["Ground"])
        self.assertEqual(ground_cdu["requestType"], "Ground Prep")

        rows, _warnings = read_spreadsheet(
            ROOT / "GILLETT - Kirke Schedule Request 10.22.26.xlsx", settings.header_synonyms
        )
        high_band = catalog.match("Tower", rows[0].sow, settings.request_types["Tower"])
        self.assertEqual(high_band["requestType"], "SNAP")

        model = SowClassifier.load(settings.path(settings.catalog_path))
        pdu_model = model.predict("Ground", "PDU upgrade", settings.request_types["Ground"])
        self.assertEqual(pdu_model["requestType"], "PDU")
        self.assertGreaterEqual(pdu_model["confidence"], 90)
        hub = model.predict(
            "Ground",
            "Installation of Nokia IXR-e routers, Samsung CDU-30 units, VDU blades, and Trimm PDUs inside the Verizon hub. Zones 20 thru 25.",
            settings.request_types["Ground"],
        )
        self.assertLess(hub["confidence"], 90)
        self.assertEqual(hub["requestType"], "")


class SampleFileTests(unittest.TestCase):
    def test_readers_find_the_real_keys(self):
        settings = load_settings(ROOT)
        gillett_rows, warnings = read_spreadsheet(ROOT / "GILLETT - Kirke Schedule Request 10.22.26.xlsx", settings.header_synonyms)
        self.assertEqual(warnings, [])
        self.assertEqual([row.fuze_id for row in gillett_rows], ["17548835", "17548835", "17548835"])
        self.assertEqual(gillett_rows[0].site_name, "GILLETT")
        self.assertEqual(gillett_rows[0].start_date.isoformat(), "2026-10-21")
        self.assertEqual(gillett_rows[1].start_date.isoformat(), "2026-10-22")
        self.assertEqual(gillett_rows[0].start_time.hour, 8)
        self.assertEqual(gillett_rows[0].end_time.hour, 18)

        powell_rows, _powell_warnings = read_spreadsheet(
            ROOT / "10-23-2026_Kirke_Schedule_Request_2026.xlsx", settings.header_synonyms
        )
        self.assertEqual(len(powell_rows), 1)
        self.assertEqual(powell_rows[0].fuze_id, "17520302")
        self.assertEqual(powell_rows[0].ps_location, "108447")
        self.assertEqual(powell_rows[0].sow, "PDU upgrade")
        self.assertEqual(powell_rows[0].start_time.hour, 23)

        gillett_mop = read_mop(ROOT / "GILLETT - MOP DAY 2.docx")
        self.assertEqual(gillett_mop.fuze_ids, [])
        self.assertEqual(gillett_mop.site_name, "GILLETT")
        self.assertIn("BBU Swap for the SNAP", gillett_mop.scope_text)

        powell_mop = read_mop(ROOT / "POWELL_PROJECT_MOP_2026.docx")
        self.assertEqual(powell_mop.fuze_ids, ["17520302"])
        self.assertEqual(powell_mop.site_name, "POWELL")
        self.assertEqual(powell_mop.job_description, "PDU Upgrade")

        jerry = read_mop(ROOT / "JERRY_BOYLES_RD_MOD_PROJECT_MOP_2026.docx")
        self.assertEqual(jerry.fuze_ids, ["17520291"])
        self.assertIn("GROUND:", jerry.scope_text)
        self.assertIn("uAPDF", jerry.scope_text)

    def test_gillett_builds_three_rows_and_two_mops(self):
        settings = load_settings(ROOT)
        with tempfile.TemporaryDirectory() as tmp:
            settings.output_dir = Path(tmp)
            result = process_package(
                settings,
                {
                    "emailId": "sample-gillett",
                    "category": "Tower",
                    "originalFrom": "Jane GC <jane.gc@vendor.com>",
                    "originalDate": "Tuesday, October 20, 2026 8:14 AM",
                    "originalTo": "tower-mop@example.com",
                    "spreadsheet": str(ROOT / "GILLETT - Kirke Schedule Request 10.22.26.xlsx"),
                    "mops": [str(ROOT / "GILLETT - MOP DAY 2.docx")],
                    "body": "",
                },
            )
            self.assertEqual(result["processingStatus"], "completed")
            self.assertEqual(result["emailStatus"], "pending_delivery")
            self.assertNotEqual(result["emailStatus"], "completed")
            self.assertEqual(result["rowCount"], 3)
            self.assertEqual(result["mopCount"], 2)
            self.assertEqual([row["requestType"] for row in result["rows"]], ["SNAP", "SNAP", "SNAP"])
            self.assertGreaterEqual(min(row["confidence"] for row in result["rows"]), 90)

            workbook = load_workbook(Path(tmp) / "Tower" / "Tower.xlsx")
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet["A1"].value, "Kirke Created Checkbox")
            self.assertEqual(sheet["F2"].value, 17548835)
            self.assertEqual(sheet["H2"].value, 107763)
            self.assertEqual(sheet["J2"].value, "136685, 1369350")
            self.assertEqual(sheet["L2"].value, "GILLETT")
            self.assertEqual(sheet["P2"].value, "KGI WIRELESS")
            self.assertEqual(sheet["S2"].value, "SNAP")
            self.assertEqual(sheet["B2"].value.date().isoformat(), "2026-10-20")
            self.assertEqual(sheet["E2"].value.date().isoformat(), "2026-10-21")
            self.assertEqual(sheet["U2"].value.date().isoformat(), "2026-10-21")
            self.assertEqual(sheet["V2"].value.hour, 8)
            self.assertEqual(sheet["W4"].value.date().isoformat(), "2026-10-22")
            self.assertEqual(sheet["AC2"].value, "jane.gc@vendor.com")
            self.assertIsNone(sheet["Y2"].value)
            self.assertIsNone(sheet["M2"].value)
            source_sow = read_spreadsheet(
                ROOT / "GILLETT - Kirke Schedule Request 10.22.26.xlsx", settings.header_synonyms
            )[0][0].sow
            self.assertEqual(sheet["Z2"].value, source_sow)
            for row in range(2, 5):
                for column in range(1, 33):
                    value = sheet.cell(row, column).value
                    self.assertFalse(isinstance(value, str) and value.startswith("="))
            workbook.close()

            source_hash = hashlib.sha256((ROOT / "GILLETT - MOP DAY 2.docx").read_bytes()).hexdigest()
            written = list((Path(tmp) / "Tower" / "mops").glob("*.docx"))
            self.assertEqual(sorted(path.name for path in written), [
                "17548835_2026-10-21.docx",
                "17548835_2026-10-22.docx",
            ])
            for path in written:
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), source_hash)

    def test_powell_without_a_category_stays_on_review(self):
        settings = load_settings(ROOT)
        with tempfile.TemporaryDirectory() as tmp:
            settings.output_dir = Path(tmp)
            result = process_package(
                settings,
                {
                    "emailId": "sample-powell",
                    "category": "Review",
                    "spreadsheet": str(ROOT / "10-23-2026_Kirke_Schedule_Request_2026.xlsx"),
                    "mops": [str(ROOT / "POWELL_PROJECT_MOP_2026.docx")],
                },
            )
            self.assertEqual(result["emailStatus"], "review_required")
            self.assertEqual(result["rows"][0]["requestType"], "")
            self.assertTrue((Path(tmp) / "Review" / "Review.xlsx").exists())
            self.assertFalse((Path(tmp) / "Ground" / "Ground.xlsx").exists())
            self.assertFalse((Path(tmp) / "Tower" / "Tower.xlsx").exists())


class MopNameTests(unittest.TestCase):
    def test_a_new_email_does_not_rename_a_delivered_file(self):
        delivered = {
            "emailId": "sample-gillett",
            "mopSource": "C:/mail/gillett.docx",
            "fuzeProjectId": "17548835",
            "startDate": "2026-10-22",
            "sourceRow": 4,
            "mopFileName": "17548835_2026-10-22.docx",
        }
        same_window = {
            "emailId": "sample-gillett",
            "mopSource": "C:/mail/gillett.docx",
            "fuzeProjectId": "17548835",
            "startDate": "2026-10-22",
            "sourceRow": 5,
            "mopFileName": "17548835_2026-10-22.docx",
        }
        newcomer = {
            "emailId": "aaa-earlier-sort",
            "mopSource": "C:/mail/other.docx",
            "fuzeProjectId": "17548835",
            "startDate": "2026-10-22",
            "sourceRow": 3,
        }
        rows = [newcomer, delivered, same_window]
        assign_mop_names(rows, "YYYY-MM-DD")
        self.assertEqual(delivered["mopFileName"], "17548835_2026-10-22.docx")
        self.assertEqual(same_window["mopFileName"], "17548835_2026-10-22.docx")
        self.assertEqual(newcomer["mopFileName"], "17548835_2026-10-22_2.docx")

    def test_a_saved_name_must_match_the_base_and_extension(self):
        backup = {
            "emailId": "old",
            "mopSource": "C:/mail/site.docx",
            "fuzeProjectId": "17548835",
            "startDate": "2026-10-22",
            "sourceRow": 2,
            "mopFileName": "17548835_2026-10-22_backup.docx",
        }
        wrong_ext = {
            "emailId": "old-doc",
            "mopSource": "C:/mail/site.docx",
            "fuzeProjectId": "17520302",
            "startDate": "2026-10-19",
            "sourceRow": 2,
            "mopFileName": "17520302_2026-10-19.doc",
        }
        assign_mop_names([backup, wrong_ext], "YYYY-MM-DD")
        self.assertEqual(backup["mopFileName"], "17548835_2026-10-22.docx")
        self.assertEqual(wrong_ext["mopFileName"], "17520302_2026-10-19.docx")


if __name__ == "__main__":
    unittest.main()
