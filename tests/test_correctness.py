import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from processor.classify_email import classify_email
from processor.export_kirke import assign_mop_names
from processor.match_mops import filename_digit_runs
from processor.pipeline import _parse_request_date, process_package
from processor.read_inputs import MopDocument, SheetRow
from processor.request_type import classify_request_type, prepare_brief_sow
from processor.settings import load_settings
from processor.match_mops import match_rows

ROOT = Path(__file__).resolve().parents[1]
TOWER = ["tower-mop@example.com"]
GROUND = ["ground-mop@example.com"]

INNER = """-----Original Message-----
From: Jane GC <jane.gc@vendor.com>
Sent: Tuesday, October 20, 2026 8:14 AM
To: Tower MOP <tower-mop@example.com>
Subject: GILLETT

Schedule attached.
"""


def test_delivered_name_stays_when_an_earlier_email_collides():
    delivered = {
        "emailId": "email-b",
        "mopSource": "C:/mail/b.docx",
        "fuzeProjectId": "17548835",
        "startDate": "2026-10-22",
        "sourceRow": 4,
    }
    assign_mop_names([delivered], "YYYY-MM-DD")
    assert delivered["mopFileName"] == "17548835_2026-10-22.docx"

    newcomer = {
        "emailId": "email-a",
        "mopSource": "C:/mail/a.docx",
        "fuzeProjectId": "17548835",
        "startDate": "2026-10-22",
        "sourceRow": 2,
    }
    assign_mop_names([newcomer, delivered], "YYYY-MM-DD")
    assert delivered["mopFileName"] == "17548835_2026-10-22.docx"
    assert newcomer["mopFileName"] == "17548835_2026-10-22_2.docx"


def test_rows_from_the_same_mop_share_one_name():
    first = _named_row("email-b", 4)
    second = _named_row("email-b", 5)
    assign_mop_names([first, second], "YYYY-MM-DD")
    assert first["mopFileName"] == second["mopFileName"] == "17548835_2026-10-22.docx"


def test_a_saved_name_is_replaced_when_the_start_date_changes():
    row = _named_row("email-b", 4)
    row["mopFileName"] = "17548835_2026-10-21.docx"
    row["startDate"] = "2026-10-22"
    assign_mop_names([row], "YYYY-MM-DD")
    assert row["mopFileName"] == "17548835_2026-10-22.docx"


def test_single_forward_uses_the_inner_to_line():
    result = classify_email({"id": "1", "body": INNER}, TOWER, GROUND)
    assert result["category"] == "Tower"
    assert result["toSource"] == "innermost"
    assert result["originalFrom"] == "jane.gc@vendor.com"
    assert "October 20, 2026" in result["originalDate"]


def test_double_forward_ignores_the_middle_hop():
    body = """-----Original Message-----
From: Outer Desk <outer@example.com>
Sent: Monday, October 19, 2026 1:00 PM
To: Ground MOP <ground-mop@example.com>
Subject: FW: job

Please see below.

From: Middle Desk <middle@example.com>
Sent: Monday, October 19, 2026 2:00 PM
To: Ground MOP <ground-mop@example.com>
Subject: FW: job

From: Jane GC <jane.gc@vendor.com>
Sent: Tuesday, October 20, 2026 8:14 AM
To: Tower MOP <tower-mop@example.com>
Subject: job

Schedule attached.
"""
    result = classify_email({"id": "2", "body": body}, TOWER, GROUND)
    assert result["category"] == "Tower"
    assert result["toSource"] == "innermost"
    assert result["originalTo"] == "tower-mop@example.com"
    assert result["originalFrom"] == "jane.gc@vendor.com"


def test_empty_innermost_to_walks_outward():
    body = """-----Original Message-----
From: Middle Desk <middle@example.com>
Sent: Monday, October 19, 2026 2:00 PM
To: Tower MOP <tower-mop@example.com>
Subject: FW: job

From: Jane GC <jane.gc@vendor.com>
Sent: Tuesday, October 20, 2026 8:14 AM
Subject: job

The original To line was dropped.
"""
    result = classify_email({"id": "3", "body": body}, TOWER, GROUND)
    assert result["category"] == "Tower"
    assert result["toSource"] == "outer_block_1"
    assert result["originalFrom"] == "middle@example.com"


def test_folded_to_line_still_parses():
    body = """-----Original Message-----
From: Jane GC <jane.gc@vendor.com>
Sent: Tuesday, October 20, 2026 8:14 AM
To: Tower MOP
 <tower-mop@example.com>
Subject: job

Schedule attached.
"""
    result = classify_email({"id": "4", "body": body}, TOWER, GROUND)
    assert result["category"] == "Tower"
    assert result["originalTo"] == "tower-mop@example.com"


def test_html_body_parses_like_plain_text():
    html = (
        "-----Original Message-----<br>"
        "<b>From:</b> Jane GC &lt;jane.gc@vendor.com&gt;<br>"
        "<b>Sent:</b> Tuesday, October 20, 2026 8:14 AM<br>"
        "<b>To:</b> Tower MOP &lt;tower-mop@example.com&gt;<br>"
        "Schedule attached."
    )
    result = classify_email(
        {"id": "5", "body": {"contentType": "html", "content": html}},
        TOWER,
        GROUND,
    )
    assert result["category"] == "Tower"
    assert result["toSource"] == "innermost"
    assert result["originalTo"] == "tower-mop@example.com"


TOWER_SENDERS = ["ctx_tower_kirke@yahoo.com"]
GROUND_SENDERS = ["ctx_ground_kirke@yahoo.com"]


def test_gmail_tower_sender_is_tower():
    result = classify_email(
        {
            "id": "gmail-tower",
            "from": "CTX Tower <ctx_tower_kirke@yahoo.com>",
            "to": "Inbox <gcbservices139@gmail.com>",
            "body": "Schedule attached.",
        },
        [],
        [],
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Tower"
    assert result["status"] == "classified"
    assert result["toSource"] == "from"
    assert result["originalFrom"] == "ctx_tower_kirke@yahoo.com"
    assert result["originalTo"] == "gcbservices139@gmail.com"


def test_gmail_ground_sender_is_ground():
    result = classify_email(
        {
            "id": "gmail-ground",
            "fromAddress": "ctx_ground_kirke@yahoo.com",
            "body": "Schedule attached.",
        },
        [],
        [],
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Ground"
    assert result["originalFrom"] == "ctx_ground_kirke@yahoo.com"


def test_gmail_body_from_line_is_used_when_the_object_has_no_sender():
    body = """From: Ground Kirke <ctx_ground_kirke@yahoo.com>
To: Inbox <gcbservices139@gmail.com>
Date: Wednesday, October 22, 2026 8:00 AM

Schedule attached.
"""
    result = classify_email(
        {"id": "gmail-body", "body": body},
        [],
        [],
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Ground"
    assert result["originalFrom"] == "ctx_ground_kirke@yahoo.com"
    assert result["originalTo"] == "gcbservices139@gmail.com"
    assert "October 22, 2026" in result["originalDate"]


def test_gmail_graph_from_shape_is_tower():
    result = classify_email(
        {
            "id": "gmail-graph",
            "from": {"emailAddress": {"address": "ctx_tower_kirke@yahoo.com"}},
            "body": "Attached.",
        },
        [],
        [],
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Tower"
    assert result["originalFrom"] == "ctx_tower_kirke@yahoo.com"


def test_gmail_unknown_sender_is_review():
    result = classify_email(
        {
            "id": "gmail-unknown",
            "sender": "other@yahoo.com",
            "to": "gcbservices139@gmail.com",
            "body": "See attached.",
        },
        TOWER,
        GROUND,
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Review"
    assert "Unrecognized sender" in result["reason"]


def test_gmail_empty_body_uses_the_sender():
    result = classify_email(
        {"id": "gmail-empty", "from": "CTX <ctx_tower_kirke@yahoo.com>", "body": "   "},
        [],
        [],
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Tower"
    assert result["originalFrom"] == "ctx_tower_kirke@yahoo.com"


def test_sender_configured_on_both_lists_is_review():
    result = classify_email(
        {"id": "gmail-both", "from": "ctx_tower_kirke@yahoo.com", "body": "hi"},
        [],
        [],
        tower_senders=["ctx_tower_kirke@yahoo.com"],
        ground_senders=["ctx_tower_kirke@yahoo.com"],
    )
    assert result["category"] == "Review"
    assert "both" in result["reason"]


def test_forward_inner_to_wins_over_the_outer_sender():
    body = """From: CTX Tower <ctx_tower_kirke@yahoo.com>
To: gcbservices139@gmail.com

-----Original Message-----
From: Jane GC <jane.gc@vendor.com>
Sent: Tuesday, October 20, 2026 8:14 AM
To: Ground MOP <ground-mop@example.com>
Subject: job

Schedule attached.
"""
    result = classify_email(
        {"id": "fwd-wins", "from": "ctx_tower_kirke@yahoo.com", "body": body},
        TOWER,
        GROUND,
        tower_senders=TOWER_SENDERS,
        ground_senders=GROUND_SENDERS,
    )
    assert result["category"] == "Ground"
    assert result["originalFrom"] == "jane.gc@vendor.com"
    assert result["originalTo"] == "ground-mop@example.com"


def test_site_name_fallback_skips_a_mop_for_another_fuze():
    row = _sheet("111", "Site X")
    foreign = _mop("foreign.docx", ["222"], "Site X")
    pairs, _unmatched = match_rows([row], [foreign], "both")
    assert pairs[0][1] is None
    assert pairs[0][2] == "missing_mop"


def test_site_name_fallback_keeps_a_mop_with_no_fuze_id():
    row = _sheet("111", "Site X")
    unlabeled = _mop("gillett.docx", [], "Site X")
    pairs, _unmatched = match_rows([row], [unlabeled], "both")
    assert pairs[0][1] is unlabeled
    assert pairs[0][2] == "site_name"


def test_unknown_match_strategy_is_rejected(tmp_path):
    config = tmp_path / "settings.json"
    config.write_text('{"matchStrategy": "fuzzy"}', encoding="utf-8")
    with pytest.raises(ValueError, match="matchStrategy"):
        load_settings(tmp_path, config)


def test_fuze_filename_match_uses_separate_digit_runs():
    stem = "MOP_12345_2026-10-20"
    runs = filename_digit_runs(stem)
    assert "12345" in runs
    assert "4520" not in runs
    assert "2026" not in runs
    assert "1234520261020" not in runs
    assert filename_digit_runs("12345") == ["12345"]


def test_ambiguous_mops_are_named_once(tmp_path):
    sheet = tmp_path / "schedule.xlsx"
    _workbook(sheet, [("123", "Vendor Co", "Site X", "cable", date(2026, 10, 21))])
    first = tmp_path / "a.docx"
    second = tmp_path / "b.docx"
    _docx(first, "123")
    _docx(second, "123")
    settings = load_settings(ROOT)
    settings.output_dir = tmp_path / "out"
    result = process_package(
        settings,
        {
            "emailId": "ambiguous",
            "category": "Review",
            "spreadsheet": str(sheet),
            "mops": [str(first), str(second)],
        },
    )
    ambiguous = [issue for issue in result["issues"] if issue["type"] == "ambiguous_mop"]
    unmatched = [issue for issue in result["issues"] if issue["type"] == "unmatched_mop"]
    assert len(ambiguous) == 1
    assert ambiguous[0]["message"] == "More than one MOP matched Fuze Project ID 123: a.docx, b.docx"
    assert unmatched == []


def test_invalid_rows_count_as_excluded_sites(tmp_path):
    sheet = tmp_path / "schedule.xlsx"
    _workbook(
        sheet,
        [
            ("111", "Vendor Co", "Site X", "snap", date(2026, 10, 21)),
            ("", "", "Orphan", "", None),
        ],
    )
    mop = tmp_path / "111.docx"
    _docx(mop, None)
    settings = load_settings(ROOT)
    settings.output_dir = tmp_path / "out"
    result = process_package(
        settings,
        {
            "emailId": "partial",
            "category": "Tower",
            "originalFrom": "jane.gc@vendor.com",
            "originalDate": "October 20, 2026",
            "spreadsheet": str(sheet),
            "mops": [str(mop)],
        },
    )
    assert result["processingStatus"] == "partial"
    assert result["excludedSites"] == 1
    assert result["invalidRows"] == 1
    assert result["includedSites"] == 1


def test_messages_use_the_configured_limits(monkeypatch):
    settings = load_settings(ROOT)
    settings.ai_api_key = "test-key"
    settings.request_type_min_confidence = 77
    monkeypatch.setattr(
        "processor.request_type._ai_classify",
        lambda *args, **kwargs: {"requestType": "", "confidence": 10, "reason": ""},
    )
    monkeypatch.setattr("processor.request_type._summarize", lambda *args, **kwargs: "word " * 10)
    classified = classify_request_type(settings, "Tower", "zzzz unique nonsense sow", "", "")
    assert classified["reason"] == "Second pass confidence is below 77"
    _text, issue = prepare_brief_sow("word " * 6, 5, settings)
    assert issue["message"] == "SOW summary was empty or still over the 5 word limit"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Monday, October 20, 2026 3:15 PM", "2026-10-20"),
        ("Oct 20, 2026", "2026-10-20"),
        ("Oct. 20, 2026", "2026-10-20"),
        ("20 October 2026", "2026-10-20"),
        ("20 Oct 2026", "2026-10-20"),
        ("Monday, 20 October 2026 3:15 PM", "2026-10-20"),
        ("2026-10-20", "2026-10-20"),
        ("2026-10-20T15:15:00", "2026-10-20"),
        ("2026-10-20T15:15:00Z", "2026-10-20"),
        ("2026-10-20T15:15:00+05:30", "2026-10-20"),
        ("10/20/2026", "2026-10-20"),
        ("10/20/26", "2026-10-20"),
        ("10/20/2026 3:15 PM", "2026-10-20"),
        ("Fri, Oct 23, 2026 at 11:30 PM", "2026-10-23"),
        ("Friday, October 23, 2026 11:30 PM", "2026-10-23"),
        ("Sent: Friday, October 23, 2026 11:30 PM", "2026-10-23"),
        ("2026-10-23T18:00:18Z", "2026-10-23"),
        ("2026-10-23T18:00:18.000Z", "2026-10-23"),
        ("", ""),
        ("not a date", ""),
    ],
)
def test_request_date_formats(raw, expected):
    assert _parse_request_date(raw) == expected


def test_body_date_is_used_when_original_date_is_empty(tmp_path):
    sheet = tmp_path / "schedule.xlsx"
    _workbook(sheet, [("111", "Vendor Co", "Site X", "snap", date(2026, 10, 21))])
    mop = tmp_path / "111.docx"
    _docx(mop, "111")
    settings = load_settings(ROOT)
    settings.output_dir = tmp_path / "out"
    result = process_package(
        settings,
        {
            "emailId": "dated",
            "category": "Tower",
            "originalFrom": "ctx_tower_kirke@yahoo.com",
            "originalDate": "",
            "body": "Date: Fri, Oct 23, 2026 at 11:30 PM\n\nSchedule attached.",
            "spreadsheet": str(sheet),
            "mops": [str(mop)],
        },
    )
    assert result["processingStatus"] == "completed"
    assert result["rowsAdded"] == 1
    assert result["outputs"]["workbook"].endswith("Tower_111_2026-10-23.xlsx")
    assert not any(issue["code"] == "DATE_MISSING" for issue in result["issues"])


def test_a_repeat_updates_the_master_instead_of_appending(tmp_path):
    sheet = tmp_path / "schedule.xlsx"
    _workbook(sheet, [("111", "Vendor Co", "Site X", "snap", date(2026, 10, 22))])
    mop = tmp_path / "111.docx"
    _docx(mop, "111")
    settings = load_settings(ROOT)
    settings.output_dir = tmp_path / "out"
    job = {
        "category": "Tower",
        "originalFrom": "ctx_tower_kirke@yahoo.com",
        "originalDate": "Fri, Oct 23, 2026 at 11:30 PM",
        "spreadsheet": str(sheet),
        "mops": [str(mop)],
    }
    first = process_package(settings, {**job, "emailId": "one"})
    second = process_package(settings, {**job, "emailId": "two", "originalFrom": "other.gc@vendor.com"})
    master = load_workbook(Path(settings.output_dir) / "Tower" / "Tower.xlsx")
    assert master.active.max_row == 2
    assert master.active["AC2"].value == "other.gc@vendor.com"
    assert master.active["B2"].value.date().isoformat() == "2026-10-23"
    master.close()
    assert first["rowsAdded"] == 1
    assert second["rowsUpdated"] == 1
    assert any(issue["code"] == "DUPLICATE_ROW" for issue in second["issues"])
    assert second["outputs"]["workbook"].endswith("Tower_111_2026-10-23.xlsx")


def test_untyped_tower_row_is_written_to_review(tmp_path):
    sheet = tmp_path / "schedule.xlsx"
    _workbook(
        sheet,
        [
            ("111", "Vendor Co", "MARLOW", "snap", date(2026, 10, 21)),
            ("222", "Vendor Co", "MARLOW", "Battery string replacement", date(2026, 10, 29)),
        ],
    )
    first = tmp_path / "111.docx"
    second = tmp_path / "222.docx"
    _docx(first, "111")
    _docx(second, "222")
    settings = load_settings(ROOT)
    settings.output_dir = tmp_path / "out"
    result = process_package(
        settings,
        {
            "emailId": "mixed",
            "category": "Tower",
            "originalFrom": "ctx_tower_kirke@yahoo.com",
            "originalDate": "Fri, Oct 23, 2026 at 11:30 PM",
            "spreadsheet": str(sheet),
            "mops": [str(first), str(second)],
        },
    )
    assert result["processingStatus"] == "review_required"
    assert [row["fuzeProjectId"] for row in result["rows"]] == ["111"]
    assert [row["fuzeProjectId"] for row in result["reviewRows"]] == ["222"]
    assert any(issue["code"] == "REQUEST_TYPE_UNMAPPED" and "MARLOW" in issue["message"] for issue in result["issues"])
    assert result["outputs"]["workbook"].endswith("Tower_111_2026-10-23.xlsx")
    assert result["outputs"]["reviewWorkbook"].endswith("Review_222_2026-10-23.xlsx")
    tower = load_workbook(Path(settings.output_dir) / "Tower" / "Tower.xlsx")
    review = load_workbook(Path(settings.output_dir) / "Review" / "Review.xlsx")
    assert tower.active.max_row == 2
    assert tower.active["F2"].value == 111
    assert review.active.max_row == 2
    assert review.active["F2"].value == 222
    assert review.active["S2"].value in (None, "")
    tower.close()
    review.close()


def _named_row(email_id: str, source_row: int) -> dict:
    return {
        "emailId": email_id,
        "mopSource": "C:/mail/b.docx",
        "fuzeProjectId": "17548835",
        "startDate": "2026-10-22",
        "sourceRow": source_row,
    }


def _sheet(fuze_id: str, site_name: str) -> SheetRow:
    return SheetRow(fuze_id, "", "", site_name, "Vendor", "sow", None, None, None, None, 2)


def _mop(name: str, fuze_ids: list[str], site_name: str) -> MopDocument:
    return MopDocument(Path(name), "", "", fuze_ids, site_name, "", "", "", "", "")


def _workbook(path: Path, rows: list[tuple]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["Fuze ID", "GC Vendor", "Site Name", "BRIEF SOW", "Start Date"])
    for row in rows:
        sheet.append(list(row))
    book.save(path)
    book.close()


def _docx(path: Path, fuze_id: str | None) -> None:
    if fuze_id:
        inner = (
            "<w:tbl><w:tr>"
            "<w:tc><w:p><w:r><w:t>Fuze ID</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>Site Name</w:t></w:r></w:p></w:tc>"
            "</w:tr><w:tr>"
            f"<w:tc><w:p><w:r><w:t>{fuze_id}</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>Site X</w:t></w:r></w:p></w:tc>"
            "</w:tr></w:tbl>"
        )
    else:
        inner = "<w:p><w:r><w:t>scope</w:t></w:r></w:p>"
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{inner}</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)
