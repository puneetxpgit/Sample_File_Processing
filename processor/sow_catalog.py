from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

_CACHE: dict[tuple[str, int], "SowCatalog"] = {}


@dataclass(frozen=True)
class CatalogExample:
    category: str
    request_type: str
    sow: str
    normalized: str
    tokens: frozenset[str]


class SowCatalog:
    def __init__(self, examples: list[CatalogExample]):
        self.examples = examples

    @classmethod
    def load(cls, path: Path) -> "SowCatalog":
        if not path.exists():
            return cls([])
        key = (str(path.resolve()), path.stat().st_mtime_ns)
        cached = _CACHE.get(key)
        if cached is not None:
            return cached
        workbook = load_workbook(path, data_only=True, read_only=True)
        examples: list[CatalogExample] = []
        for category in ("Tower", "Ground"):
            if category not in workbook.sheetnames:
                continue
            sheet = workbook[category]
            for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
                request_type = str(row[0] or "").strip()
                sow = row[2] if isinstance(row[2], str) else str(row[2] or "")
                if not request_type or request_type.casefold() in {"n/a", "na"}:
                    continue
                normalized = normalize_sow(sow)
                tokens = frozenset(_tokens(normalized))
                if not normalized or not tokens:
                    continue
                examples.append(
                    CatalogExample(category, request_type, re.sub(r"\s+", " ", sow).strip(), normalized, tokens)
                )
        workbook.close()
        catalog = cls(examples)
        _CACHE[key] = catalog
        return catalog

    def match(self, category: str, text: str, allowed: list[str]) -> dict | None:
        normalized = normalize_sow(text)
        if not normalized:
            return None
        pool = [
            example
            for example in self.examples
            if example.category == category and example.request_type in allowed
        ]
        exact = [example for example in pool if example.normalized == normalized]
        exact_hit = _single_type(exact)
        if exact_hit is not None:
            return _hit(exact_hit, 98, "Exact SOW match in the category catalog")
        if len({example.request_type for example in exact}) > 1:
            return None

        incoming_tokens = frozenset(_tokens(normalized))
        contained = [
            example
            for example in pool
            if len(example.tokens) >= 8 and example.normalized in normalized
        ]
        contained_hit = _single_type(contained)
        if contained_hit is not None:
            return _hit(contained_hit, 95, "SOW contains a catalog example for this type")
        if len({example.request_type for example in contained}) > 1:
            return None

        scored: list[tuple[float, CatalogExample]] = []
        for example in pool:
            if len(example.tokens) < 8 or not incoming_tokens:
                continue
            overlap = len(example.tokens & incoming_tokens)
            if overlap < 8:
                continue
            contain = overlap / min(len(example.tokens), len(incoming_tokens))
            jaccard = overlap / len(example.tokens | incoming_tokens)
            scored.append((0.75 * contain + 0.25 * jaccard, example))
        if not scored:
            return None
        best_score = max(score for score, _example in scored)
        if best_score < 0.84:
            return None
        leaders = [example for score, example in scored if score >= best_score - 0.02]
        leader = _single_type(leaders)
        if leader is None:
            return None
        rivals = [
            example.request_type
            for score, example in scored
            if example.request_type != leader.request_type and score >= best_score - 0.1
        ]
        if rivals:
            return None
        confidence = min(94, round(best_score * 100))
        if confidence < 90:
            return None
        return _hit(leader, confidence, "SOW is closest to this catalog example")


def normalize_sow(text: str) -> str:
    cleaned = html.unescape(str(text or ""))
    cleaned = cleaned.replace("\u2013", " ").replace("\u2014", " ").replace("\xa0", " ")
    cleaned = cleaned.casefold().replace("uapdf", "uadpf")
    cleaned = re.sub(r"[^a-z0-9]+", " ", cleaned)
    return re.sub(r" +", " ", cleaned).strip()


def _tokens(normalized: str) -> list[str]:
    return [token for token in normalized.split() if len(token) > 1]


def _single_type(examples: list[CatalogExample]) -> CatalogExample | None:
    types = {example.request_type for example in examples}
    if len(types) != 1:
        return None
    return max(examples, key=lambda example: len(example.tokens))


def _hit(example: CatalogExample, confidence: int, reason: str) -> dict:
    preview = example.sow if len(example.sow) <= 140 else example.sow[:137] + "..."
    return {
        "requestType": example.request_type,
        "confidence": confidence,
        "stage": "catalog",
        "reason": f"{reason}: {preview}",
    }
