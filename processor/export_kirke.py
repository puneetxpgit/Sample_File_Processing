from __future__ import annotations

import json
import re
import shutil
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from processor.settings import Settings

FMT_REQUEST = "mm-dd-yy"
FMT_START = "mm-dd-yy"
FMT_DAY = 'm"/"d'
FMT_TIME = 'h":"mm" "AM/PM'

COLUMNS = {
    "request_date": 2,
    "start_date_e": 5,
    "fuze": 6,
    "ps_location": 8,
    "enodeb": 10,
    "site_name": 12,
    "gc_vendor": 16,
    "request_type": 19,
    "start_date_u": 21,
    "start_time": 22,
    "end_date": 23,
    "end_time": 24,
    "brief_sow": 26,
    "gc_email": 29,
}


def category_dir(settings: Settings, category: str) -> Path:
    path = settings.path(settings.output_dir) / category
    path.mkdir(parents=True, exist_ok=True)
    return path


def manifest_dir(settings: Settings, category: str) -> Path:
    path = category_dir(settings, category) / "manifests"
    path.mkdir(parents=True, exist_ok=True)
    return path


def write_manifest(settings: Settings, category: str, email_id: str, manifest: dict) -> Path:
    path = manifest_dir(settings, category) / f"{safe_id(email_id)}.json"
    path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return path


def reset_category(settings: Settings, category: str) -> None:
    folder = settings.path(settings.output_dir) / category
    if folder.exists():
        shutil.rmtree(folder)


def schedule_key(row: dict) -> tuple:
    return (
        str(row.get("fuzeProjectId") or ""),
        str(row.get("startDate") or ""),
        str(row.get("startTime") or ""),
        str(row.get("briefSow") or ""),
    )


def schedule_sort_key(row: dict) -> tuple:
    return (
        str(row.get("startDate") or ""),
        str(row.get("startTime") or ""),
        str(row.get("fuzeProjectId") or ""),
    )


def reply_workbook_name(category: str, email_id: str, rows: list[dict]) -> str:
    fuzes = list(dict.fromkeys(str(row.get("fuzeProjectId") or "") for row in rows if row.get("fuzeProjectId")))
    request_dates = [str(row.get("requestDate")) for row in rows if row.get("requestDate")]
    if request_dates:
        stamp = request_dates[0]
    else:
        starts = sorted(str(row.get("startDate")) for row in rows if row.get("startDate"))
        stamp = starts[0] if starts else "UNDATED"
    stem = fuzes[0] if len(fuzes) == 1 else safe_id(email_id)
    return f"{category}_{_safe_stem(stem)}_{stamp}.xlsx"


def rebuild_category(settings: Settings, category: str) -> dict:
    folder = category_dir(settings, category)
    manifest_paths = sorted((folder / "manifests").glob("*.json"))
    manifests = [(path, json.loads(path.read_text(encoding="utf-8"))) for path in manifest_paths]
    workbook_path = folder / f"{category}.xlsx"
    mop_dir = folder / "mops"
    mop_dir.mkdir(parents=True, exist_ok=True)

    ranked = []
    for index, (path, manifest) in enumerate(manifests):
        mtime = path.stat().st_mtime_ns
        for row in manifest.get("rows") or []:
            row["emailId"] = manifest.get("emailId")
            ranked.append((mtime, index, row))
    rows = _dedupe_rows(ranked)
    assign_mop_names(rows, settings.mop_date_format)
    for path, manifest in manifests:
        path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    rows.sort(key=schedule_sort_key)
    _write_workbook(settings, workbook_path, rows)

    expected = {}
    for row in rows:
        source = row.get("mopSource")
        dest_name = row.get("mopFileName")
        if source and dest_name:
            expected[dest_name] = source
    for existing in mop_dir.iterdir():
        if existing.is_file() and existing.name not in expected:
            existing.unlink()
    for dest_name, source in expected.items():
        target = mop_dir / dest_name
        if not target.exists() or target.read_bytes() != Path(source).read_bytes():
            shutil.copyfile(source, target)
    return {
        "workbook": str(workbook_path),
        "rowCount": len(rows),
        "mopCount": len(expected),
        "mops": [str(mop_dir / name) for name in sorted(expected)],
    }


def write_reply(settings: Settings, category: str, email_id: str, rows: list[dict]) -> dict:
    """Write this email's workbook and MOP copies without names inherited from older runs."""
    folder = category_dir(settings, category) / "replies"
    folder.mkdir(parents=True, exist_ok=True)
    assign_reply_mop_names(rows, settings.mop_date_format)
    ordered = sorted(rows, key=schedule_sort_key)
    workbook_path = folder / reply_workbook_name(category, email_id, ordered)
    _write_workbook(settings, workbook_path, ordered)
    expected = {}
    for row in ordered:
        source = row.get("mopSource")
        dest_name = row.get("mopFileName")
        if source and dest_name:
            expected[dest_name] = source
    for dest_name, source in expected.items():
        shutil.copyfile(source, folder / dest_name)
    return {
        "workbook": str(workbook_path),
        "mops": [str(folder / name) for name in sorted(expected)],
    }


def assign_reply_mop_names(rows: list[dict], date_format: str) -> None:
    """Name this email's MOP copies from Fuze ID and start date only."""
    used: set[str] = set()
    assigned: dict[tuple[str, str], str] = {}
    for row in rows:
        source = row.get("mopSource") or ""
        ext = Path(source).suffix if source else ".docx"
        stamp = format_stamp(row.get("startDate") or "", date_format)
        base = f"{_safe_stem(row.get('fuzeProjectId') or 'NO_FUZE')}_{stamp}"
        key = (source, base)
        if key in assigned:
            row["mopFileName"] = assigned[key]
            continue
        name = f"{base}{ext}"
        suffix = 2
        while name.casefold() in used:
            name = f"{base}_{suffix}{ext}"
            suffix += 1
        used.add(name.casefold())
        assigned[key] = name
        row["mopFileName"] = name


def _dedupe_rows(ranked: list[tuple[int, int, dict]]) -> list[dict]:
    """Keep one row per Fuze ID, start date, start time, and BRIEF SOW.

    A later manifest replaces the earlier row so a repeat updates the master.
    """
    best: dict[tuple, tuple[tuple[int, int], dict]] = {}
    for mtime, index, row in ranked:
        key = schedule_key(row)
        rank = (mtime, index)
        current = best.get(key)
        if current is None or rank >= current[0]:
            best[key] = (rank, row)
    return [item[1] for item in best.values()]


def assign_mop_names(rows: list[dict], date_format: str) -> None:
    """Name MOP files without shifting names that were already delivered.

    A later email can sort ahead of an older one. Rebuilding from scratch would
    give the new row the plain filename and rename the delivered file to _2.
    Saved names that still match this row's Fuze ID, date, and extension stay put.
    """
    used: set[str] = set()
    assigned: dict[tuple[str, str], str] = {}
    ordered = sorted(rows, key=lambda row: (row.get("emailId") or "", row.get("startDate") or "", row.get("sourceRow") or 0))

    def base_and_ext(row: dict) -> tuple[str, str, str]:
        source = row.get("mopSource") or ""
        ext = Path(source).suffix if source else ".docx"
        stamp = format_stamp(row.get("startDate") or "", date_format)
        base = f"{_safe_stem(row.get('fuzeProjectId') or 'NO_FUZE')}_{stamp}"
        return source, base, ext

    pending = []
    for row in ordered:
        source, base, ext = base_and_ext(row)
        existing = row.get("mopFileName") or ""
        key = (source, base)
        if key in assigned:
            row["mopFileName"] = assigned[key]
        elif _name_fits(existing, base, ext) and existing.casefold() not in used:
            used.add(existing.casefold())
            assigned[key] = existing
        else:
            pending.append(row)

    for row in pending:
        source, base, ext = base_and_ext(row)
        key = (source, base)
        if key in assigned:
            row["mopFileName"] = assigned[key]
            continue
        name = f"{base}{ext}"
        suffix = 2
        while name.casefold() in used:
            name = f"{base}_{suffix}{ext}"
            suffix += 1
        used.add(name.casefold())
        assigned[key] = name
        row["mopFileName"] = name


def _name_fits(existing: str, base: str, ext: str) -> bool:
    if not existing or not ext:
        return False
    return re.fullmatch(
        rf"{re.escape(base)}(?:_\d+)?{re.escape(ext)}",
        existing,
        re.IGNORECASE,
    ) is not None


def format_stamp(iso_date: str, date_format: str) -> str:
    if not iso_date:
        return "UNDATED"
    try:
        parsed = datetime.strptime(iso_date, "%Y-%m-%d").date()
    except ValueError:
        return "UNDATED"
    return parsed.strftime(_strftime_format(date_format))


def _write_workbook(settings: Settings, workbook_path: Path, rows: list[dict]) -> None:
    template = settings.path(settings.template_path)
    workbook = load_workbook(template)
    sheet = workbook["Sheet3"] if "Sheet3" in workbook.sheetnames else workbook.active
    for name in list(workbook.sheetnames):
        if workbook[name] is not sheet:
            del workbook[name]
    workbook.active = sheet
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for offset, row in enumerate(rows, start=2):
        _write_row(sheet, offset, row, settings.populate_location_and_site_name)
    workbook.save(workbook_path)
    workbook.close()


def _write_row(sheet, excel_row: int, row: dict, include_location_and_site: bool) -> None:
    _set_date(sheet, excel_row, COLUMNS["request_date"], row.get("requestDate"), FMT_REQUEST)
    _set_date(sheet, excel_row, COLUMNS["start_date_e"], row.get("startDate"), FMT_START)
    _set_id(sheet, excel_row, COLUMNS["fuze"], row.get("fuzeProjectId"))
    if include_location_and_site:
        _set_id(sheet, excel_row, COLUMNS["ps_location"], row.get("psLocation"))
        _set_text(sheet, excel_row, COLUMNS["site_name"], row.get("siteName"))
    _set_text(sheet, excel_row, COLUMNS["enodeb"], row.get("enodeb"))
    _set_text(sheet, excel_row, COLUMNS["gc_vendor"], row.get("gcVendor"))
    _set_text(sheet, excel_row, COLUMNS["request_type"], row.get("requestType"))
    _set_date(sheet, excel_row, COLUMNS["start_date_u"], row.get("startDate"), FMT_DAY)
    _set_time(sheet, excel_row, COLUMNS["start_time"], row.get("startTime"))
    _set_date(sheet, excel_row, COLUMNS["end_date"], row.get("endDate"), FMT_DAY)
    _set_time(sheet, excel_row, COLUMNS["end_time"], row.get("endTime"))
    cell = _set_text(sheet, excel_row, COLUMNS["brief_sow"], row.get("briefSow"))
    cell.alignment = Alignment(wrap_text=True)
    _set_text(sheet, excel_row, COLUMNS["gc_email"], row.get("gcEmail"))


def _set_text(sheet, excel_row: int, column: int, value) -> object:
    cell = sheet.cell(excel_row, column, None if value in (None, "") else value)
    return cell


def _set_id(sheet, excel_row: int, column: int, value) -> None:
    text = "" if value is None else str(value).strip()
    if text.isdigit():
        sheet.cell(excel_row, column, int(text))
    elif text:
        sheet.cell(excel_row, column, text)


def _set_date(sheet, excel_row: int, column: int, iso_value: str | None, number_format: str) -> None:
    cell = sheet.cell(excel_row, column, None)
    cell.number_format = number_format
    if not iso_value:
        return
    parsed = datetime.strptime(iso_value, "%Y-%m-%d")
    cell.value = parsed


def _set_time(sheet, excel_row: int, column: int, value: str | None) -> None:
    cell = sheet.cell(excel_row, column, None)
    cell.number_format = FMT_TIME
    if not value:
        return
    parsed = datetime.strptime(value, "%H:%M:%S").time()
    cell.value = parsed


def _strftime_format(pattern: str) -> str:
    return pattern.replace("YYYY", "%Y").replace("DD", "%d").replace("MM", "%m")


def _safe_stem(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in str(value))
    return cleaned.strip("_") or "NO_FUZE"


def safe_id(email_id: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in email_id)
    return cleaned[:80] or "email"


def iso_date(value: date | None) -> str:
    return value.isoformat() if value else ""


def iso_time(value: time | None) -> str:
    return value.strftime("%H:%M:%S") if value else ""
