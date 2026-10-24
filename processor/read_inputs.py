from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class SheetRow:
    fuze_id: str
    ps_location: str
    enodeb: str
    site_name: str
    gc_vendor: str
    sow: str
    start_date: date | None
    end_date: date | None
    start_time: time | None
    end_time: time | None
    source_row: int
    warnings: list[str] = field(default_factory=list)


@dataclass
class MopDocument:
    path: Path
    text: str
    header_text: str
    fuze_ids: list[str]
    site_name: str
    ps_location: str
    site_number: str
    job_description: str
    scope_text: str
    date_window: str
    read_error: str = ""


def read_spreadsheet(path: Path, synonyms: dict[str, list[str]]) -> tuple[list[SheetRow], list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        grid = _xlsx_grid(path)
    elif suffix == ".csv":
        grid = _csv_grid(path)
    elif suffix == ".xls":
        grid = _xls_grid(path)
    else:
        raise ValueError(f"Unsupported spreadsheet type: {path.name}")

    header_index, columns = _find_header(grid, synonyms)
    if header_index is None:
        raise ValueError(f"No header row with a Fuze ID column was found in {path.name}")

    rows: list[SheetRow] = []
    warnings: list[str] = []
    for offset, raw in enumerate(grid[header_index + 1 :], start=header_index + 2):
        mapped = {name: _cell(raw, col) for name, col in columns.items()}
        if not any(_present(value) for value in mapped.values()):
            continue
        fuze = _as_id(mapped.get("Fuze Project ID"))
        vendor = str(mapped.get("GC Vendor") or "").strip()
        if not fuze or not vendor:
            missing = [name for name, present in (("Fuze Project ID", fuze), ("GC Vendor", vendor)) if not present]
            warnings.append(
                f"Row {offset} excluded. Missing required field: {', '.join(missing)}"
            )
            continue
        start_date, start_warn = _as_date(mapped.get("Start Date"), f"Row {offset} Start Date")
        end_date, end_warn = _as_date(mapped.get("End Date"), f"Row {offset} End Date")
        start_time, start_time_warn = _as_time(mapped.get("Start Time"), f"Row {offset} Start Time")
        end_time, end_time_warn = _as_time(mapped.get("End Time"), f"Row {offset} End Time")
        row_warnings = [item for item in (start_warn, end_warn, start_time_warn, end_time_warn) if item]
        rows.append(
            SheetRow(
                fuze_id=fuze,
                ps_location=_as_id(mapped.get("PS Location")),
                enodeb=_as_text(mapped.get("EnodeB")),
                site_name=str(mapped.get("Site Name") or "").strip(),
                gc_vendor=vendor,
                sow=mapped.get("SOW") if isinstance(mapped.get("SOW"), str) else _as_text(mapped.get("SOW")),
                start_date=start_date,
                end_date=end_date,
                start_time=start_time,
                end_time=end_time,
                source_row=offset,
                warnings=row_warnings,
            )
        )
    return rows, warnings


def read_mop(path: Path) -> MopDocument:
    if path.suffix.lower() != ".docx":
        raise ValueError(
            f"{path.name} could not be read. Legacy {path.suffix} needs conversion to .docx before processing."
        )
    try:
        paragraphs, table_rows, header_text = _docx_parts(path)
    except Exception as exc:
        raise ValueError(f"Failed to read {path.name}: {exc}") from exc

    text = "\n".join(paragraphs)
    site_name = ""
    fuze_ids: list[str] = []
    ps_location = ""
    site_number = ""
    date_window = ""
    for index, row in enumerate(table_rows):
        headers = [cell.strip().casefold() for cell in row]
        if "fuze id" not in headers or index + 1 >= len(table_rows):
            continue
        values = table_rows[index + 1]
        if not values or "fuze id" in {cell.strip().casefold() for cell in values}:
            continue
        data = {headers[i]: values[i].strip() for i in range(min(len(headers), len(values)))}
        fuze = re.sub(r"\D", "", data.get("fuze id", ""))
        if fuze:
            fuze_ids.append(fuze)
        site_name = site_name or data.get("site name", "")
        ps_location = ps_location or re.sub(r"\s+", "", data.get("ps loc #", ""))
        site_number = site_number or re.sub(r"\s+", "", data.get("site #", ""))
        date_window = date_window or data.get("date", "")

    header_site = _site_from_header(header_text)
    if header_site and not site_name:
        site_name = header_site

    job = _section_value(paragraphs, "JOB DESCRIPTION")
    if re.search(r"scope of work\s*:", text, re.I):
        scope = re.split(r"scope of work\s*:", text, maxsplit=1, flags=re.I)[1].strip()
    else:
        scope = text
    return MopDocument(
        path=path,
        text=text,
        header_text=header_text,
        fuze_ids=list(dict.fromkeys(fuze_ids)),
        site_name=site_name.strip(),
        ps_location=ps_location,
        site_number=site_number,
        job_description=job,
        scope_text=scope,
        date_window=date_window,
    )


def _docx_parts(path: Path) -> tuple[list[str], list[list[str]], str]:
    paragraphs: list[str] = []
    table_rows: list[list[str]] = []
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        body = root.find(f"{W}body")
        for child in list(body):
            tag = child.tag.split("}")[-1]
            if tag == "p":
                line = _xml_text(child).strip()
                if line:
                    paragraphs.append(line)
            elif tag == "tbl":
                for tr in child.findall(f".//{W}tr"):
                    cells = []
                    for tc in tr.findall(f"{W}tc"):
                        bits = [_xml_text(p).strip() for p in tc.findall(f"{W}p")]
                        bits = [bit for bit in bits if bit]
                        if bits:
                            cells.append("".join(bits) if all(re.fullmatch(r"\d+", bit) for bit in bits) else " ".join(bits))
                    if len(cells) >= 2:
                        table_rows.append(cells)
        headers = []
        for name in archive.namelist():
            if name.startswith("word/header") and name.endswith(".xml"):
                headers.append(_xml_text(ET.fromstring(archive.read(name))).strip())
    return paragraphs, table_rows, "\n".join(bit for bit in headers if bit)


def _xml_text(element) -> str:
    return "".join((node.text or "") for node in element.iter(f"{W}t"))


def _site_from_header(header_text: str) -> str:
    match = re.search(r"SITE INFO AND SCHEDULE[_\s-]+(.+)", header_text or "", re.I)
    return match.group(1).strip(" _-") if match else ""


def _section_value(paragraphs: list[str], heading: str) -> str:
    target = heading.casefold()
    for index, line in enumerate(paragraphs):
        if line.strip().casefold() == target and index + 1 < len(paragraphs):
            return paragraphs[index + 1].strip()
    return ""


def _xlsx_grid(path: Path) -> list[list]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook.active
    grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return grid


def _csv_grid(path: Path) -> list[list]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [row for row in csv.reader(handle)]


def _xls_grid(path: Path) -> list[list]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError(f"Cannot read {path.name}. Install xlrd to process .xls files.") from exc
    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    grid = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            else:
                row.append(cell.value)
        grid.append(row)
    return grid


def _find_header(grid: list[list], synonyms: dict[str, list[str]]) -> tuple[int | None, dict[str, int]]:
    lookup = {}
    for canonical, names in synonyms.items():
        for name in names:
            lookup[_norm_header(name)] = canonical
    for index, row in enumerate(grid[:15]):
        columns = {}
        for col, value in enumerate(row):
            canonical = lookup.get(_norm_header(value))
            if canonical and canonical not in columns:
                columns[canonical] = col
        if "Fuze Project ID" in columns:
            return index, columns
    return None, {}


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip().casefold()


def _cell(row: list, index: int):
    if index is None or index >= len(row):
        return None
    return row[index]


def _present(value) -> bool:
    return value is not None and str(value).strip() != ""


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _as_id(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _as_date(value, label: str) -> tuple[date | None, str]:
    if value is None or value == "":
        return None, ""
    if isinstance(value, datetime):
        return value.date(), ""
    if isinstance(value, date):
        return value, ""
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%m-%d-%y"):
        try:
            return datetime.strptime(text, fmt).date(), ""
        except ValueError:
            continue
    return None, f"{label} could not be parsed and was left blank: {text}"


def _as_time(value, label: str) -> tuple[time | None, str]:
    if value is None or value == "":
        return None, ""
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0), ""
    if isinstance(value, time):
        return value.replace(microsecond=0), ""
    text = str(value).strip().upper().replace(" ", "")
    match = re.fullmatch(r"(\d{1,2})(?::(\d{2}))?(AM|PM)", text)
    if not match:
        return None, f"{label} could not be parsed and was left blank: {value}"
    hour = int(match.group(1))
    minute = int(match.group(2) or 0)
    if match.group(3) == "PM" and hour != 12:
        hour += 12
    if match.group(3) == "AM" and hour == 12:
        hour = 0
    if hour > 23 or minute > 59:
        return None, f"{label} could not be parsed and was left blank: {value}"
    return time(hour, minute), ""
